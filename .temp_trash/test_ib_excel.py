import io
import sys
sys.path.insert(0, "/Users/aavibharucha/Documents/octavian")

from mna_model_engine import MnAAssumptions, get_mna_engine
from lbo_model_engine import LBOAssumptions, get_lbo_engine
from ib_excel_engine import build_mna_workbook, build_lbo_workbook, build_comps_workbook

m = get_mna_engine().run_mna(MnAAssumptions(
    acquirer_ticker="ACQ", target_ticker="TGT",
    acquirer_price=100, acquirer_eps=5, acquirer_shares=1000,
    target_price=50, target_eps=2, target_shares=200,
    offer_premium=0.30, percent_stock=0.5, percent_cash=0.5,
    cost_of_debt=0.06, tax_rate=0.21, pre_tax_synergies=100,
))
l = get_lbo_engine().run_lbo(LBOAssumptions(
    ticker="TGT", target_name="Target", entry_year=2026, exit_year=2031,
    ltm_ebitda=500, entry_multiple=10, exit_multiple=10,
    leverage_multiple=5, interest_rate=0.08, tax_rate=0.25,
))

for name, data in [("MNA", build_mna_workbook(m)), ("LBO", build_lbo_workbook(l))]:
    assert data and len(data) > 5000, f"{name} workbook too small"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    print(f"{name} workbook OK: sheets={wb.sheetnames}")
    formula_count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    print(f"  live formulas: {formula_count}")

# Comps workbook
import pandas as pd
comps = pd.DataFrame({
    "Company": ["AAA", "BBB", "CCC"],
    "Market Cap": [1e9, 2e9, 3e9],
    "P/E": [15.0, 20.0, 25.0],
    "EV/EBITDA": [8.0, 9.0, 11.0],
    "Profit Margin": [0.15, 0.20, 0.25],
})
med = pd.DataFrame({"Company": ["Median", "Mean"], "Market Cap": [2e9, 2e9],
                    "P/E": [20.0, 20.0], "EV/EBITDA": [9.0, 9.333],
                    "Profit Margin": [0.20, 0.20]})
cdata = build_comps_workbook(comps, med)
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(cdata))
print(f"COMPS workbook OK: sheets={wb.sheetnames}")
print("ALL EXCEL BUILDS OK")
