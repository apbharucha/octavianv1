"""Verify the DCF pitchbook uses real assumption values (regression for the
'decks show $5,000M / 8% / 20% regardless of company' bug)."""
import os
import sys
import io

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from financial_model_generator import DCFAssumptions, get_dcf_engine
from presentation_generator import get_presentation_generator

# Company with DISTINCTIVE inputs that would expose the old fabricated
# fallbacks ($5,000M revenue / 8% growth / 20% margin / 1,000M shares).
a = DCFAssumptions(
    ticker="ACME",
    base_revenue=12345.0,
    revenue_growth_rates=[0.14] * 5,
    ebit_margin=0.32,
    tax_rate=0.19,
    da_pct_revenue=0.04,
    capex_pct_revenue=0.06,
    nwc_change_pct_revenue=0.02,
    equity_value_market=25000.0,
    debt_value=3000.0,
    cost_of_debt=0.045,
    risk_free_rate=0.0425,
    equity_risk_premium=0.055,
    beta=1.15,
    terminal_growth_rate=0.028,
    cash=1200.0,
    shares_outstanding=850.0,
    current_price=88.0,
)
r = get_dcf_engine().run_dcf(a)
print(f"FV/share: {r.fair_value_per_share:.2f} | assumptions.base_revenue={r.assumptions.base_revenue}")

prs = get_presentation_generator()
data = prs.generate_dcf_pitchbook("ACME", r)
print(f"pptx bytes: {len(data)}")
assert len(data) > 5000, "deck should be populated"

# Extract slide text from the pptx to assert real values are present.
from pptx import Presentation
p = Presentation(io.BytesIO(data))
all_text = []
for slide in p.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_text.append(shape.text_frame.text)
text = "\n".join(all_text)

print("---- sample slide text ----")
for line in text.splitlines()[:40]:
    print(" ", line)

assert "12,345" in text, "deck must show real base revenue $12,345M"
assert "14%" in text, "deck must show real growth ~14%"
assert "32.0%" in text, "deck must show real EBIT margin 32%"
assert "2.80%" in text, "deck must show real terminal growth 2.80%"
assert "850" in text, "deck must show real shares 850M"
assert "$62.80B" in text or "$62.8" in text, "EV must be displayed in $B correctly (was $0.00B bug)"
assert "$61.00B" in text or "$61.0" in text, "Equity must be displayed in $B correctly"
assert "5000" not in text.replace("12,345", ""), "old $5,000M fallback must be gone"
print("DCF PITCHBOOK ACCURACY: PASS")
