import sys, io, re
sys.path.insert(0, "/Users/aavibharucha/Documents/octavian")
from lbo_model_engine import LBOAssumptions, get_lbo_engine
from ib_excel_engine import build_lbo_workbook
from openpyxl import load_workbook

l = get_lbo_engine().run_lbo(LBOAssumptions(
    ticker='TGT', target_name='Target', entry_year=2026, exit_year=2031,
    ltm_ebitda=500, entry_multiple=10, exit_multiple=10,
    leverage_multiple=5, interest_rate=0.08, tax_rate=0.25,
))
print("engine exit_equity:", l.exit_equity_value)
wb = load_workbook(io.BytesIO(build_lbo_workbook(l)))
for ws in wb.worksheets:
    print(f"--- {ws.title} ---")
    for row in ws.iter_rows(max_col=8):
        vals = [c.coordinate + "=" + repr(c.value) for c in row if c.value is not None]
        if vals:
            print(" | ".join(vals))
