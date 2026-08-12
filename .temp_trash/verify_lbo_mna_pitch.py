"""Verify LBO + M&A pitchbooks render real model outputs, not fabricated
45/15/40% sources and uses or hardcoded $500M synergies."""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from lbo_model_engine import LBOAssumptions, get_lbo_engine
from mna_model_engine import MnAAssumptions, get_mna_engine
from presentation_generator import get_presentation_generator
from pptx import Presentation


def _text_of(data):
    p = Presentation(io.BytesIO(data))
    parts = []
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


# ── LBO with DISTINCTIVE inputs ─────────────────────────────────────────────
la = LBOAssumptions(
    ticker="TGTX", target_name="TGTX", entry_year=2026, exit_year=2031,
    ltm_ebitda=500.0, entry_multiple=9.0, exit_multiple=10.5,
    leverage_multiple=4.5, interest_rate=0.075, tax_rate=0.25,
    revenue_growth_rates=[0.05, 0.06, 0.07, 0.06, 0.05],
    ebitda_margins=[0.22, 0.23, 0.24, 0.24, 0.25],
)
lr = get_lbo_engine().run_lbo(la)
print(f"LBO purchase={lr.purchase_price:.0f} debt={lr.debt_amount:.0f} equity={lr.equity_amount:.0f} "
      f"irr={lr.irr:.1%} moic={lr.moic:.2f}")
ld = get_presentation_generator().generate_lbo_pitchbook("TGTX", lr)
lt = _text_of(ld)
assert f"{lr.purchase_price:,.0f}" in lt, "purchase price must be real"
assert f"{lr.irr:.1%}" in lt, "IRR must be real"
assert f"{lr.moic:.2f}x" in lt, "MOIC must be real"
# fabricated 45% / 15% / 40% splits must be gone
assert "45%" not in lt or "40%" not in lt, "fabricated 45/15/40 sources must be gone"
# The debt_amount (4.5x * 500 = 2250) should appear, not pp*0.35 fabrication
assert f"{lr.debt_amount:,.0f}" in lt, "debt amount must be real"
print("LBO PITCHBOOK: PASS")


# ── M&A with distinctive inputs ─────────────────────────────────────────────
ma = MnAAssumptions(
    acquirer_ticker="ACQ", target_ticker="TGT",
    acquirer_price=150.0, acquirer_eps=8.0, acquirer_shares=2000.0,
    target_price=40.0, target_eps=2.5, target_shares=300.0,
    offer_premium=0.25, percent_stock=0.60, percent_cash=0.40,
    pre_tax_synergies=320.0, cost_of_debt=0.055, tax_rate=0.21,
    acquirer_revenue=30000.0, target_revenue=4500.0,
    acquirer_ebitda=6000.0, target_ebitda=900.0,
)
mr = get_mna_engine().run_mna(ma)
print(f"M&A offer={mr.offer_price:.2f} acc/dil={mr.accretion_dilution_pct:.1%} "
      f"syn={ma.pre_tax_synergies:.0f} deal={mr.total_deal_value:.0f}")
md = get_presentation_generator().generate_mna_pitchbook("ACQ", "TGT", mr)
mt = _text_of(md)
assert f"{mr.offer_price:.2f}" in mt, "offer price must be real"
assert "$320" in mt, "synergies must be the real 320, not 500"
assert "25%" in mt, "premium must be real 25%"
assert f"{mr.total_deal_value:,.0f}" in mt, "deal value must be real"
assert "$500" not in mt, "fabricated $500M synergy fallback must be gone"
print("M&A PITCHBOOK: PASS")
print("ALL PITCHBOOK ACCURACY TESTS PASSED")
