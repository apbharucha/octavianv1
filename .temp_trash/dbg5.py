import sys, io, re
sys.path.insert(0, "/Users/aavibharucha/Documents/octavian")
from lbo_model_engine import LBOAssumptions, get_lbo_engine
from ib_excel_engine import build_lbo_workbook
from openpyxl import load_workbook

l = get_lbo_engine().run_lbo(LBOAssumptions(
    ticker='TGT', target_name='Target', entry_year=2026, exit_year=2031,
    ltm_ebitda=500, entry_multiple=10, exit_multiple=10,
    leverage_multiple=5, interest_rate=0.08, tax_rate=0.25,
))
print('engine exit_equity:', l.exit_equity_value)
wb = load_workbook(io.BytesIO(build_lbo_workbook(l)))

_wb_sheets = {}
for ws in wb.worksheets:
    m = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                m[(cell.column_letter, cell.row)] = cell.value
    _wb_sheets[ws.title] = m

_depth = [0]
def resolve(ref, sheet):
    if _depth[0] > 25:
        return None
    if '!' in ref:
        sname, addr = ref.split('!', 1)
        sheet = _wb_sheets.get(sname)
        if sheet is None:
            return None
    else:
        addr = ref
    mm = re.match(r'^([A-Z]+)(\d+)$', addr)
    if not mm:
        return None
    cell = sheet.get((mm.group(1), int(mm.group(2))))
    if cell is None:
        return None
    if isinstance(cell, str) and cell.startswith('='):
        _depth[0] += 1
        try:
            return eval_formula(cell, sheet)
        finally:
            _depth[0] -= 1
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        return float(cell)
    try:
        return float(cell)
    except (TypeError, ValueError):
        return None

def v(token, sheet):
    if isinstance(token, (int, float)):
        return float(token)
    if re.match(r'^[A-Z]{1,3}\d+$', str(token)):
        return float(resolve(token, sheet) or 0.0)
    return float(token)

def eval_formula(formula, sheet):
    body = formula[1:].strip()
    body = re.sub(r'MIN\(([^,]+),([^)]+)\)',
                  lambda m: f'min({v(m.group(1), sheet)}, {v(m.group(2), sheet)})', body)
    body = re.sub(r'MAX\(([^,]+),([^)]+)\)',
                  lambda m: f'max({v(m.group(1), sheet)}, {v(m.group(2), sheet)})', body)
    def repl(m):
        val = resolve(m.group(0), sheet)
        return str(val) if val is not None else '0'
    body = re.sub(r'([A-Z]{1,3}\d+)', repl, body)
    try:
        return float(eval(body, {'__builtins__': {}}, {}))
    except Exception as e:
        return float('nan')

for (col, row), val in _wb_sheets['Returns'].items():
    if isinstance(val, str) and val.startswith('='):
        res = eval_formula(val, _wb_sheets['Returns'])
        print(f'{col}{row}: {val} -> {res}')
