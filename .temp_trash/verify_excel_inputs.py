"""Verify spreadsheet generators use user-entered numbers exactly."""
import io
import sys
import types
from unittest.mock import patch, MagicMock

sys.path.insert(0, "/Users/aavibharucha/Documents/octavian")

import pandas as pd
import numpy as np
import streamlit as st

import spreadsheet_generator as sg

# ── Capture streamlit download_button bytes ──────────────────────────────────
captured = {}

def fake_download(label, data, file_name, mime=None, **kw):
    captured[label] = (bytes(data), file_name)
    return None

def fake_show(*a, **k):
    return None

for fn in ["download_button", "success", "error", "info", "warning", "spinner", "dataframe", "markdown", "caption"]:
    setattr(sg.st, fn, fake_show if fn != "download_button" else fake_download)
    # spinner used as context manager
    if fn == "spinner":
        class _Spin:
            def __enter__(self): return self
            def __exit__(self, *a): return False
        setattr(sg.st, fn, lambda *a, **k: _Spin())

# ── Fake yfinance ────────────────────────────────────────────────────────────
class FakeTicker:
    def __init__(self, sym):
        self.sym = sym
        self.info = {
            "totalRevenue": 40_000_000_000,
            "ebitda": 9_000_000_000,
            "sharesOutstanding": 1_000_000_000,
            "currentPrice": 150.0,
            "totalCash": 4_000_000_000,
            "totalDebt": 8_000_000_000,
            "enterpriseValue": 48_000_000_000,
        }
        cols = pd.to_datetime(["2025-12-31", "2024-12-31", "2023-12-31"])
        self.financials = pd.DataFrame(
            {
                "Total Revenue": [40e9, 36e9, 32e9],
                "EBIT": [9e9, 8e9, 7e9],
            },
            index=cols,
        ).T
        self.balance_sheet = pd.DataFrame({c: [0] for c in cols})
        self.cashflow = pd.DataFrame({c: [0] for c in cols})

fake_yf = MagicMock()
fake_yf.Ticker.side_effect = lambda s: FakeTicker(s)
sys.modules["yfinance"] = fake_yf

from openpyxl import load_workbook

def cell_scan(wb):
    """Return {sheet: {cell: value}} for all sheets."""
    out = {}
    for ws in wb.worksheets:
        d = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = c.value
        out[ws.title] = d
    return out

failures = []

# ── TEST 1: DCF template uses user WACC/TGR/tax exactly ─────────────────────
captured.clear()
WACC, TGR, TAX = 0.105, 0.0325, 0.27
sg._generate_dcf_model("AAPL", wacc=WACC, tgr=TGR, tax_rate=TAX)
assert captured, "DCF did not emit a workbook"
data, fname = list(captured.values())[0]
wb = load_workbook(io.BytesIO(data))
cells = cell_scan(wb)
summary = cells["Valuation Output"]
# WACC and TGR must appear verbatim
vals = [v for v in summary.values() if isinstance(v, (int, float))]
if not any(abs(float(v) - WACC) < 1e-9 for v in vals):
    failures.append(f"DCF: WACC {WACC} not found verbatim in summary {list(summary.values())}")
if not any(abs(float(v) - TGR) < 1e-9 for v in vals):
    failures.append(f"DCF: TGR {TGR} not found verbatim in summary {list(summary.values())}")
# Tax rate: the Taxes row must equal EBIT * TAX (uses user rate, not 21%)
# Year-1 EBIT = rev_base*(1+14% fade start) * margin where margin starts at ebit_base/rev_base + 0.5%
dcf_sheet = cells["DCF Model"]
tax_cell = None
tax_label = None
for coord, v in dcf_sheet.items():
    if isinstance(v, str) and "Taxes" in v:
        tax_label = v
        col = "".join(ch for ch in coord if ch.isalpha())
        nxt = chr(ord(col[0]) + 1) + coord[1:]
        tax_cell = dcf_sheet.get(nxt)
        break
if tax_cell is None:
    failures.append("DCF: no Taxes row found")
elif "21%" in tax_label:
    failures.append(f"DCF: Taxes label hardcodes 21% despite tax rate {TAX}: {tax_label!r}")
else:
    # Year 1 EBIT in the model: rev=40e9, g=14% (fade start), margin=9/40+0.005=0.23
    exp_ebit = 40e9 * 1.14 * (9 / 40 + 0.005)
    if abs(float(tax_cell) + exp_ebit * TAX) > exp_ebit * TAX * 0.01:
        failures.append(f"DCF: Taxes value {tax_cell} != Year1 EBIT*user_tax ({exp_ebit*TAX})")
print(f"DCF OK  -> {fname} | label={tax_label!r} taxes={tax_cell} (expected ~{exp_ebit*TAX:.0f})")

# ── TEST 2: LBO template uses user debt%/exit mult/rate exactly ─────────────
captured.clear()
DEBT_PCT, EXIT_M, RATE = 0.65, 11.0, 0.07
sg._generate_lbo_model("AAPL", debt_pct=DEBT_PCT, exit_mult=EXIT_M, interest_rate=RATE)
assert captured, "LBO did not emit a workbook"
data, fname = list(captured.values())[0]
wb = load_workbook(io.BytesIO(data))
cells = cell_scan(wb)
su = cells["Sources & Uses"]
su_vals = list(su.values())
if not any("65%" in str(v) for v in su_vals):
    failures.append(f"LBO: debt split 65% not labeled in Sources & Uses: {su_vals}")
# Sponsor equity must be 35% of total uses
ret = cells["Returns Analysis"]
ret_vals = list(ret.values())
if not any(str(v).startswith("11.0x") for v in ret_vals):
    failures.append(f"LBO: exit multiple 11.0x not verbatim: {ret_vals}")
# Interest: Year-1 interest expense must be debt_amt * RATE
proj = cells["Projections"]
proj_rows = {}
for coord, v in proj.items():
    if isinstance(v, str):
        proj_rows[v] = coord
if "(-) Interest Exp" in proj_rows:
    coord = proj_rows["(-) Interest Exp"]
    col = "".join(ch for ch in coord if ch.isalpha())
    row = "".join(ch for ch in coord if ch.isdigit())
    y1 = chr(ord(col) + 1) + row  # Year 1 sits one column right of the label
    # expected = 65% of (EV + 2% fees) * 7%
    ev = 48_000_000_000
    total_uses = ev * 1.02
    exp_int = total_uses * DEBT_PCT * RATE
    actual = proj.get(y1)
    if actual is None or abs(float(actual) + exp_int) > 1e6:
        failures.append(f"LBO: Year-1 interest {actual} != expected {exp_int} (cell {y1})")
    print(f"LBO OK  -> {fname} | Y1 interest={actual} expected={exp_int:.0f} (cell {y1})")
else:
    failures.append("LBO: no Interest row in Projections")

# ── TEST 3: Advanced custom generator honors symbol + formatting ─────────────
captured.clear()
fake_df = pd.DataFrame({
    "Date": pd.date_range("2024-01-01", periods=60),
    "Open": np.linspace(100, 120, 60),
    "High": np.linspace(102, 122, 60),
    "Low": np.linspace(98, 118, 60),
    "Close": np.linspace(100, 121, 60),
    "Volume": np.full(60, 1_000_000),
})
with patch.object(sg, "_safe_fetch_data", return_value=fake_df), \
     patch.object(sg.st, "info", fake_show):
    spec = sg.SpreadsheetSpec(
        symbols=["MSFT", "NVDA"],
        timeframe="1y", interval="1d",
        data_columns=["Date", "Close", "Volume", "SMA_20", "RSI_14", "Returns_1d"],
        price_columns=["Close"], technical_columns=["SMA_20", "RSI_14"],
        risk_columns=["Volatility_20d"],
        chart_types=["line"], formatting={"number_format": "decimal_2", "decimal_places": 2,
                                          "conditional_formatting": True, "include_index": True,
                                          "freeze_header": True, "autofit": True, "filters": True},
        color_scheme="professional", include_summary=True, include_statistics=True,
        include_charts=True, include_correlations=True, include_financial_model=False,
        model_type=None, custom_calculations={"risk": {"sharpe": True}},
        sheet_names={"data": "Market Data", "summary": "Summary", "charts": "Charts",
                     "analysis": "Analysis", "correlation": "Correlations", "custom": "Custom"},
    )
    sg._generate_advanced_spreadsheet(spec)
assert captured, "Advanced generator did not emit a workbook"
data, fname = list(captured.values())[0]
wb = load_workbook(io.BytesIO(data))
sheets = [ws.title for ws in wb.worksheets]
for need in ["Market Data", "Summary", "Correlations"]:
    if need not in sheets:
        failures.append(f"Advanced: missing sheet {need}; got {sheets}")
# Data must contain both symbols
data_ws = wb["Market Data"]
sym_col = [c.coordinate for c in data_ws[3] if c.value == "Symbol"]
vals = []
if sym_col:
    col = sym_col[0][0]
    vals = [data_ws[f"{col}{r}"].value for r in range(4, data_ws.max_row + 1)]
if "MSFT" not in vals or "NVDA" not in vals:
    failures.append(f"Advanced: symbols MSFT/NVDA missing from data sheet: {vals}")
print(f"Advanced OK -> {fname} | sheets={sheets} | symbols={set(v for v in vals if v)}")

# ── REPORT ───────────────────────────────────────────────────────────────────
print("\n==================== RESULT ====================")
if failures:
    print("FAILURES:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("ALL EXCEL INPUT-ACCURACY TESTS PASSED")
