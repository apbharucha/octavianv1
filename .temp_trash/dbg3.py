import sys, io
sys.path.insert(0, "/Users/aavibharucha/Documents/octavian")
from mna_model_engine import MnAAssumptions, get_mna_engine
from lbo_model_engine import LBOAssumptions, get_lbo_engine
from ib_excel_engine import build_mna_workbook, build_lbo_workbook
from openpyxl import load_workbook
import re

m = get_mna_engine().run_mna(MnAAssumptions(
    acquirer_ticker='ACQ', target_ticker='TGT',
    acquirer_price=100, acquirer_eps=5, acquirer_shares=1000,
    target_price=50, target_eps=2, target_shares=200,
    offer_premium=0.30, percent_stock=0.5, percent_cash=0.5,
    cost_of_debt=0.06, tax_rate=0.21, pre_tax_synergies=100,
))
print("pf_eps:", m.pro_forma_eps, "acc_dollar:", m.accretion_dilution_dollar)
wb = load_workbook(io.BytesIO(build_mna_workbook(m)))
ws = wb["Assumptions"]
for row in ws.iter_rows(min_row=40, max_row=78, max_col=3):
    vals = [c.coordinate + "=" + repr(c.value) for c in row if c.value is not None]
    if vals:
        print(" | ".join(vals))
