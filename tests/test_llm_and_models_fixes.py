"""
Regression tests for the reported bug cluster:

1. LLM entity resolution — CAPEX / FCF / GPU / CUDA / MOAT / MULTI / GROSS /
   BUILD / FACTS / ASPS / LOSS / UNIT / CYCLE / EV / FIVE ... must NEVER
   resolve to tickers; real tickers (NVDA, AAPL) must; ambiguous tokens (AI,
   BASE, MOAT) resolve ONLY when referenced as securities.
2. Mega-query decomposition — a single deep-dive thesis must NOT explode into
   dozens of repetitive parts; genuine multi-part queries decompose into at
   most 6 sections with unique labels.
3. Breaking trades — zero/NaN prices produce NO setup (never $0 targets), and
   alerts/notes are asset-specific and include real price levels.
4. DCF auto-fill — no NoneType crash when yfinance `info` is missing, and
   revenue is derived correctly (never the per-share fallback that produced
   ~$100M instead of ~$390B for AAPL).
5. Pitchbook accuracy — DCF/M&A/LBO decks read REAL model fields instead of
   fabricated fallbacks ($5,000M revenue, 8% growth, $0.00B EV, ...).
"""
import io
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from financial_llm_engine import (  # noqa: E402
    expand_query_intents,
    _decompose_mega_query,
)


# ─────────────────────────────────────────────────────────────────────────────
# 1. ENTITY RESOLUTION
# ─────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("concept", [
    "CAPEX", "FCF", "DCF", "EBITDA", "EBIT", "EPS", "NOPAT", "ROIC", "ROCE",
    "CAGR", "WACC", "RSI", "MACD", "NIM", "MARGIN", "REVENUE", "GROWTH",
    "YIELD", "DIVIDEND", "PAYOUT", "LEVERAGE", "LIQUIDITY", "SOLVENCY",
    "BETA", "ALPHA", "SHARPE", "SORTINO", "VOL", "VOLATILITY", "DRAWDOWN",
    "PREMIUM", "DISCOUNT", "SPREAD", "NOTIONAL", "EXPIRY", "STRIKE", "IV",
    "IMPLIED", "THETA", "GAMMA", "VEGA", "DELTA", "RHO", "VANNA", "CHARM",
    "VOLGA", "GPU", "CPU", "CUDA", "ASIC", "TPU", "NPU", "FPGA",
    "INFERENCE", "TRAINING", "LLM", "MODEL", "DATACENTER", "HYPESCALER",
    "NTM", "TTM", "FWD", "YTD", "QTD", "MOM", "YOY", "FY24", "FY25",
    "Q1", "Q2", "Q3", "Q4", "H1", "H2", "CONSENSUS", "ESTIMATE", "GUIDANCE",
    "BEAT", "MISS", "MULTI", "GROSS", "BUILD", "FACTS", "ASPS", "LOSS",
    "UNIT", "CYCLE", "RATE", "SPEED", "PACE", "RANGE", "STAGE", "SCALE",
    "SCOPE", "TOTAL", "CORE", "TOP", "BOTTOM", "END", "START",
    "MID", "HIGH", "WIDE", "NARROW", "FAST", "SLOW", "BIG", "SMALL", "LARGE",
    "MEDIUM", "FULL",
])
def test_concept_words_never_resolve_to_tickers(concept):
    """Semantic/finance concepts must never appear as tickers, even in ALL
    CAPS inside a NVDA-style deep-dive prompt."""
    q = (f"Analyze NVDA and its {concept} trends over the next 12 months. "
         f"Compare {concept} with AMD.")
    _intents, tickers, _sectors = expand_query_intents(q)
    assert concept not in tickers, f"{concept} leaked as ticker in {tickers}"
    assert "NVDA" in tickers
    assert "AMD" in tickers


def test_reported_nvda_prompt_produces_only_real_tickers():
    """The exact reported bug: the NVDA deep-dive prompt must yield ONLY
    ['NVDA', 'AMD'] — no CAPEX / FCF / GPU / CUDA / MOAT / EV / MULTI / ..."""
    q = ("Conduct a full fundamental, quantitative, macroeconomic and sentiment-driven "
         "investment assessment of NVIDIA (NVDA). Reconstruct revenue and earnings "
         "drivers by segment. Evaluate competitive position against AMD, custom ASICs "
         "and hyperscaler silicon. Determine whether CUDA represents a durable moat. "
         "Build a DCF using at least three explicit scenarios. Perform a reverse DCF. "
         "Compare using EV/Revenue, EV/EBITDA, P/E and FCF yield. Analyze sensitivity "
         "to Fed policy, real interest rates, USD strength and the semiconductor cycle. "
         "Construct at least five scenarios and assign probabilities. Identify the 10 "
         "strongest arguments. Give me probability-weighted fair value and the "
         "probability of a >30% drawdown.")
    _intents, tickers, _sectors = expand_query_intents(q)
    assert set(tickers) == {"NVDA", "AMD"}, tickers


def test_ratio_shorthand_never_yields_metric_ticker():
    for ratio in ["EV/EBITDA", "EV/Revenue", "EV/EBIT", "EV/FCF", "P/E", "P / E"]:
        _intents, tickers, _sectors = expand_query_intents(
            f"What is the {ratio} for NVDA?")
        metric = ratio.split("/")[0].split()[0].strip()
        assert metric not in tickers, f"{metric} leaked from {ratio}: {tickers}"
        assert "NVDA" in tickers


@pytest.mark.parametrize("num_word", [
    "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE",
    "TEN", "TWENTY", "THIRTY", "FORTY", "FIFTY", "HUNDRED", "THOUSAND",
    "MILLION", "BILLION", "FIRST", "SECOND", "THIRD", "HALF",
])
def test_number_words_never_resolve(num_word):
    _intents, tickers, _sectors = expand_query_intents(
        f"Estimate {num_word} scenarios for NVDA")
    assert num_word not in tickers, f"{num_word} leaked: {tickers}"
    assert "NVDA" in tickers


@pytest.mark.parametrize("tok,sec_ctx,con_ctx", [
    ("AI", ["AI stock", "AI shares", "AI price", "buy AI", "outlook for AI"],
     ["AI infrastructure", "AI capex", "AI demand", "AI spending"]),
    ("BASE", ["BASE stock", "BASE earnings", "outlook for BASE"],
     ["base case", "base scenario", "base rate"]),
    ("MOAT", ["MOAT ETF", "buy MOAT"], ["durable moat", "competitive moat"]),
])
def test_ambiguous_tokens_need_security_context(tok, sec_ctx, con_ctx):
    for ctx in sec_ctx:
        _intents, tickers, _sectors = expand_query_intents(
            f"{ctx} — compare with NVDA")
        assert tok in tickers, f"'{ctx}' should resolve {tok} as a security: {tickers}"
    for ctx in con_ctx:
        _intents, tickers, _sectors = expand_query_intents(ctx)
        assert tok not in tickers, f"'{ctx}' is a concept, not {tok}: {tickers}"


def test_loss_and_unit_never_resolve():
    """LOSS / UNIT are NOT real securities in the universe — even "LOSS stock"
    must not produce a phantom ticker (the universe gate is final)."""
    for ctx in ("LOSS stock", "net loss", "operating loss", "LOSS",
                "unit economics", "UNIT", "unit growth"):
        _intents, tickers, _sectors = expand_query_intents(ctx)
        assert "LOSS" not in tickers, f"'{ctx}' leaked LOSS: {tickers}"
        assert "UNIT" not in tickers, f"'{ctx}' leaked UNIT: {tickers}"


@pytest.mark.parametrize("tok,sec_ctx,con_ctx", [
    ("LOW", ["LOW stock", "outlook for LOW", "buy LOW", "LOW shares"],
     ["low volatility", "low volume", "low growth", "low margin"]),
    ("KEY", ["KEY stock", "outlook for KEY"], ["key drivers", "key risks", "key catalysts"]),
    ("MAIN", ["MAIN stock", "outlook for MAIN"], ["main driver", "main risk"]),
])
def test_real_tickers_that_look_like_english(tok, sec_ctx, con_ctx):
    """LOW (Lowe's), KEY (KeyCorp), MAIN (Main Street Capital) are REAL tickers
    — they must resolve as securities but NEVER as English adjectives/nouns."""
    for ctx in sec_ctx:
        _intents, tickers, _sectors = expand_query_intents(
            f"{ctx} — compare with NVDA")
        assert tok in tickers, f"'{ctx}' should resolve {tok}: {tickers}"
    for ctx in con_ctx:
        _intents, tickers, _sectors = expand_query_intents(ctx)
        assert tok not in tickers, f"'{ctx}' is English, not {tok}: {tickers}"


# ─────────────────────────────────────────────────────────────────────────────
# 2. MEGA-QUERY DECOMPOSITION
# ─────────────────────────────────────────────────────────────────────────────
def test_deep_dive_thesis_not_exploded_into_parts():
    """The reported 35-part bug: a single deep-dive NVDA thesis is ONE request
    and must not be decomposed into repetitive mini-sections."""
    q = ("Conduct a full fundamental assessment of NVIDIA (NVDA) over the next "
         "12-24 months. Reconstruct revenue drivers by segment. Evaluate "
         "competitive position against AMD, custom ASICs and hyperscaler silicon. "
         "Determine whether CUDA represents a durable moat. Build a DCF using at "
         "least three explicit scenarios. Perform a reverse DCF. Compare using "
         "EV/Revenue, EV/EBITDA, P/E and FCF yield. Analyze sensitivity to Fed "
         "policy and real rates. Construct a comprehensive risk matrix. Identify "
         "catalysts over 0-3, 3-6, 6-12 and 12-24 months. Construct at least five "
         "scenarios and assign probabilities that sum to 100%. Identify the 10 "
         "strongest arguments against your conclusion. Give me probability-weighted "
         "fair value, 12-month expected return, and the probability of a >30% "
         "drawdown.")
    parts = _decompose_mega_query(q)
    assert parts is None, f"deep-dive thesis must stay single, got {len(parts)} parts"


def test_multi_part_query_decomposes_within_cap():
    """A genuine multi-part query (outlook + trade setup + probability) must
    decompose into at most 6 parts with distinct ticker targets."""
    q = ("Give me an outlook for AAPL, then a trade setup for MSFT, and what is "
         "the probability NVDA pulls back 25%")
    parts = _decompose_mega_query(q)
    assert parts is not None
    assert 2 <= len(parts) <= 6, f"expected 2-6 parts, got {len(parts)}: {parts}"


def test_comma_joined_multi_ask_decomposes():
    """Comma-joined multi-asks targeting DIFFERENT instruments must split;
    deep-dives about ONE company must not."""
    q = "Analyze MSFT earnings, hedge my AMD position with puts, and compare TSLA vs NFLX on growth"
    parts = _decompose_mega_query(q)
    assert parts is not None and 2 <= len(parts) <= 6, parts

    # Same pattern but all fragments about the SAME company → single thesis
    single = ("Analyze Apple (AAPL): reconstruct the revenue model by product "
              "and geography, estimate services growth and gross margin "
              "trajectory, build a three-stage DCF, compare against EV/EBITDA "
              "comps, run a reverse DCF on the current price, stress-test under "
              "rising rates, list the top risks and catalysts, and give a final "
              "probability-weighted fair value with confidence levels.")
    parts_single = _decompose_mega_query(single)
    assert parts_single is None, f"single-company deep-dive must not split: {parts_single}"


# ─────────────────────────────────────────────────────────────────────────────
# 3. BREAKING TRADES — price accuracy + dynamic alerts
# ─────────────────────────────────────────────────────────────────────────────
class _FakeSignal:
    def __init__(self, direction="BULLISH", confidence=0.65, expected_return=0.05):
        self.direction = direction
        self.confidence = confidence
        self.expected_return = expected_return


def _bt_gen():
    import breaking_trades_generator as btg

    gen = btg.BreakingTradesGenerator.__new__(btg.BreakingTradesGenerator)
    gen.min_confidence = 0.0
    gen._effective_min_confidence = 0.0
    gen.quant = None
    gen._cache = {}
    return gen


def test_breaking_trades_zero_price_never_emits_setup():
    """The reported bug: suggested assets priced at $0 with $0 targets. A
    0/NaN current price must produce NO price levels (caller skips symbol)."""
    gen = _bt_gen()
    ind = {"atr": 2.0, "rsi": 55, "momentum_10d": 4, "momentum_20d": 6,
           "volume_ratio": 1.4, "volatility": 0.22, "trend_strength": 0.6}
    assert gen._calculate_price_levels(0.0, "BULLISH", ind, {}) is None
    assert gen._calculate_price_levels(float("nan"), "BULLISH", ind, {}) is None
    assert gen._calculate_price_levels(float("inf"), "BULLISH", ind, {}) is None
    assert gen._calculate_price_levels(-5.0, "BULLISH", ind, {}) is None


def test_breaking_trades_levels_are_positive_and_ordered():
    gen = _bt_gen()
    ind = {"atr": 2.0, "rsi": 55, "momentum_10d": 4, "momentum_20d": 6,
           "volume_ratio": 1.4, "volatility": 0.22, "trend_strength": 0.6}
    lv = gen._calculate_price_levels(100.0, "BULLISH", ind, {})
    assert lv is not None
    assert lv["entry"] > 0 and lv["stop_loss"] > 0
    assert lv["entry"] > lv["stop_loss"]
    assert lv["tp1"] > lv["entry"] > lv["stop_loss"]
    assert lv["risk_reward"] >= 1.0

    lv_bear = gen._calculate_price_levels(100.0, "BEARISH", ind, {})
    assert lv_bear is not None
    assert lv_bear["entry"] < lv_bear["stop_loss"]
    assert lv_bear["tp1"] < lv_bear["entry"] < lv_bear["stop_loss"]


def test_breaking_trades_reasoning_is_asset_specific_with_real_levels():
    gen = _bt_gen()
    ind = {"atr": 2.0, "rsi": 55, "momentum_10d": 4, "momentum_20d": 6,
           "volume_ratio": 1.4, "volatility": 0.22, "trend_strength": 0.6}
    levels = gen._calculate_price_levels(100.0, "BULLISH", ind, {})
    sig = _FakeSignal(confidence=0.65)
    scores = {"score": 80, "confidence_score": 0.65}
    r = gen._generate_reasoning("NVDA", sig, ind, {"type": "Momentum Breakout"},
                                scores, levels=levels, confidence_score=0.65)
    # Alerts must carry the real levels, not generic text
    joined = " ".join(r.get("alerts", []))
    assert "NVDA" in r.get("primary", "")
    assert f"${levels['entry']:.2f}" in joined or f"${levels['entry']}" in joined
    assert f"${levels['stop_loss']:.2f}" in joined or f"${levels['stop_loss']}" in joined
    assert "$0.00" not in joined


# ─────────────────────────────────────────────────────────────────────────────
# 4. DCF AUTO-FILL — NoneType crash + revenue derivation
# ─────────────────────────────────────────────────────────────────────────────
def _make_financials(**rows):
    """yfinance-style financials DataFrame: rows = line items, cols = dates."""
    return pd.DataFrame({"2026-06-30": rows})


def _fake_ticker(info=None, financials=None, price=150.0, shares=1e9):
    import types

    # NOTE: class bodies do NOT close over enclosing function locals for names
    # that are also assigned in the body ("info = info" -> NameError). Bind to
    # distinct names first.
    _info, _financials, _price, _shares = info, financials, price, shares

    class FakeTicker:
        info = _info
        financials = _financials

        @property
        def fast_info(self):
            return types.SimpleNamespace(last_price=_price, shares=_shares)

        def history(self, period="1mo", **kw):
            idx = pd.date_range("2026-06-01", periods=10, freq="D")
            return pd.DataFrame({"Close": np.linspace(_price * 0.97, _price, 10)},
                                index=idx)

        def get_info(self):
            return _info or {}

    return FakeTicker()


def _no_network(monkeypatch):
    """Offline-safe: never hit get_realtime_price during auto-fill tests."""
    monkeypatch.setattr("data_sources.get_realtime_price", lambda *a, **k: (None, None))


def test_autofill_no_none_crash_when_info_missing(monkeypatch):
    """Regression: t.info == None (a known yfinance failure mode) previously
    raised 'argument of type NoneType is not iterable' inside the auto-fill."""
    import financial_model_generator as fmg
    import yfinance as yf

    _no_network(monkeypatch)
    fin = _make_financials(**{"Net Income": 100_000_000_000.0})
    fake = _fake_ticker(info=None, financials=fin, price=223.0,
                        shares=15_500_000_000)
    monkeypatch.setattr(yf, "Ticker", lambda sym: fake)
    data = fmg.fetch_ticker_fundamentals("AAPL")
    assert isinstance(data, dict)
    assert data.get("ticker") == "AAPL"


def test_autofill_derives_real_revenue_scale(monkeypatch):
    """Regression: when totalRevenue was missing, the per-share fallback
    (~$25 for AAPL) failed the >1e6 check and clamped revenue to $100M
    instead of ~$390B. Revenue must be derived at the correct scale."""
    import financial_model_generator as fmg

    _no_network(monkeypatch)
    fin = _make_financials(**{
        "Total Revenue": 390_000_000_000.0,
        "EBITDA": 130_000_000_000.0,
        "Net Income": 100_000_000_000.0,
    })
    info = {"sharesOutstanding": 15_500_000_000, "currentPrice": 223.0,
            "marketCap": 3_450_000_000_000.0, "totalRevenue": 390_000_000_000.0}
    fake = _fake_ticker(info=info, financials=fin, price=223.0,
                        shares=15_500_000_000)
    import yfinance as yf

    monkeypatch.setattr(yf, "Ticker", lambda sym: fake)
    data = fmg.fetch_ticker_fundamentals("AAPL")
    assert abs(data.get("revenue_m", 0) - 390_000) < 2_000, (
        f"revenue should be ~$390,000M, got {data.get('revenue_m')}")
    assert data.get("ebitda_m", 0) > 100_000
    assert data.get("market_cap_m", 0) > 2_000_000


def test_autofill_does_not_fabricate_ebitda(monkeypatch):
    """Regression: the old _extract_ebitda_millions returned a fabricated
    500.0 fallback when EBITDA was missing. Missing data must be None/absent,
    never an invented number."""
    import financial_model_generator as fmg

    _no_network(monkeypatch)
    fin = _make_financials(**{"Net Income": 10_000_000_000.0})  # no EBITDA row
    info = {"sharesOutstanding": 1e9, "currentPrice": 50.0,
            "marketCap": 50_000_000_000.0, "totalRevenue": 50_000_000_000.0}
    fake = _fake_ticker(info=info, financials=fin, price=50.0, shares=1e9)
    import yfinance as yf

    monkeypatch.setattr(yf, "Ticker", lambda sym: fake)
    data = fmg.fetch_ticker_fundamentals("AAPL")
    assert data.get("ebitda_m") in (None, 0), (
        f"EBITDA must not be fabricated, got {data.get('ebitda_m')}")


# ─────────────────────────────────────────────────────────────────────────────
# 5. PITCHBOOK ACCURACY — real model fields, not fabricated fallbacks
# ─────────────────────────────────────────────────────────────────────────────
def test_dcf_pitchbook_shows_real_values():
    """The DCF deck previously showed hardcoded $5,000M revenue / 8% growth /
    20% EBIT margin regardless of the model. It must show the ACTUAL inputs
    and real $B enterprise value (not $0.00B from a /1e9 unit bug)."""
    import financial_model_generator as fmg
    from presentation_generator import get_presentation_generator

    a = fmg.DCFAssumptions(
        ticker="ACME", base_revenue=12345.0, revenue_growth_rates=[0.14] * 5,
        ebit_margin=0.32, tax_rate=0.19, da_pct_revenue=0.04,
        capex_pct_revenue=0.06, nwc_change_pct_revenue=0.02,
        equity_value_market=25000.0, debt_value=3000.0, cost_of_debt=0.045,
        risk_free_rate=0.0425, equity_risk_premium=0.055, beta=1.15,
        terminal_growth_rate=0.028, cash=1200.0, shares_outstanding=850.0,
        current_price=88.0,
    )
    r = fmg.get_dcf_engine().run_dcf(a)
    assert r.enterprise_value > 0

    b = get_presentation_generator().generate_dcf_pitchbook("ACME", r)
    assert b and b.startswith(b"PK")
    from pptx import Presentation

    prs = Presentation(io.BytesIO(b))
    text = ""
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                text += sh.text_frame.text + "\n"
    assert "$12,345M" in text or "12,345" in text, "deck must show real base revenue"
    assert "~14%" in text, "deck must show real 14% growth"
    assert "32.0%" in text, "deck must show real 32% EBIT margin"
    assert "$0.00B" not in text, "EV must be displayed in correct $B units"
    assert "$5,000M" not in text, "fabricated fallback revenue must be gone"
    assert "~8%" not in text, "fabricated 8% growth must be gone"


# ─────────────────────────────────────────────────────────────────────────────
# 6. ROUTING FIXES — deep-dive priority, TGT context, sector/FX/valuation
# ─────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("tok,sec_ctx", [
    ("TGT", ["growth stocks like TGT", "key catalysts for TGT",
              "portfolio with TGT, BKR and XLB", "such as TGT and LOW",
              "upside and downside risks for TGT"]),
    ("COST", ["stocks like COST", "names like COST", "catalysts for COST"]),
    ("LOW", ["stocks like LOW", "catalysts for LOW", "with LOW and HD"]),
])
def test_ambiguous_tickers_resolve_in_stock_list_context(tok, sec_ctx):
    """The reported medium-bucket failures: TGT / COST / LOW typed in list /
    catalyst / portfolio contexts ('growth stocks like TGT', 'key catalysts
    for TGT', 'portfolio with TGT, BKR and XLB') must resolve as securities."""
    for ctx in sec_ctx:
        _intents, tickers, _sectors = expand_query_intents(ctx)
        assert tok in tickers, f"'{ctx}' should resolve {tok}: {tickers}"


def test_banks_plural_maps_to_financials_sector():
    """'Top banks stocks for growth' must detect the financials sector via the
    plural alias 'banks' -> 'bank' (previously the singular-only match missed
    it, so the query fell to the generic macro template)."""
    _intents, tickers, sectors = expand_query_intents("Top banks stocks for growth")
    assert "financials" in sectors, sectors


def test_fx_intent_flagged_without_pair_name():
    """'Which currency is the best buy right now?' is an FX question even
    though no explicit pair is named — the majors must be injected."""
    _intents, tickers, _sectors = expand_query_intents("Which currency is the best buy right now?")
    assert _intents.get("fx") is True
    assert any(str(t).endswith("=X") for t in tickers), tickers


def test_valuation_intent_flagged():
    _intents, tickers, _sectors = expand_query_intents(
        "Assess the valuation of LOW using DCF logic: what growth and margin assumptions "
        "does the current price already embed?")
    assert _intents.get("valuation") is True
    assert "LOW" in tickers


def test_deep_dive_memo_stays_single_despite_numbered_sections():
    """The LOW memo regression: a 'write a complete investment memo on X with
    sections (1)-(6)' request is ONE coherent thesis. It must NOT be
    decomposed into generic per-section mini-analyses."""
    q = ("Write a complete investment memo on LOW for the investment committee. "
         "Sections: (1) thesis; (2) fundamental drivers by segment (growth, margins, "
         "cash generation); (3) valuation with DCF and comps; (4) risks; "
         "(5) catalysts; (6) final recommendation.")
    _intents, tickers, _sectors = expand_query_intents(q)
    assert "LOW" in tickers
    from financial_llm_engine import _is_deep_dive_query
    assert _is_deep_dive_query(q, tickers) is True


@pytest.mark.parametrize("q,expect_fx", [
    ("Which currency is the best buy right now?", True),
    ("Is the dollar strengthening?", True),
    ("Analyze NFLX's competitive moat versus FXI", False),
    ("What are the key catalysts for FXI over the next 6 months?", False),
    ("Compare DIS and FXI on growth", False),
])
def test_fx_intent_never_fires_on_fxi(q, expect_fx):
    """'fx' is a substring of the ETF ticker FXI — the FX intent must match
    on word boundaries, never inside 'FXI'."""
    intents, _t, _s = expand_query_intents(q)
    assert bool(intents.get("fx")) is expect_fx, (q, intents)


def test_pe_substring_never_fires_valuation_on_xpev():
    """'pe' is a substring of XPEV — valuation must not fire on the ticker."""
    intents, tickers, _s = expand_query_intents("How does XPEV compare to UAL?")
    assert "XPEV" in tickers and "UAL" in tickers
    assert intents.get("valuation") is False, intents
    assert intents.get("comparison") is True


def test_commodity_ask_keeps_commodities_intent():
    """'Is heating oil a good buy?' is a commodity question (HO=F), not a
    sector scan of energy stocks."""
    intents, tickers, sectors = expand_query_intents("Is heating oil a good buy right now?")
    assert intents.get("commodities") is True
    assert any(str(t).endswith("=F") for t in tickers), tickers


def test_undervalued_tech_keeps_valuation_intent():
    """'Which technology stocks are undervalued?' is a valuation screening
    question, not a momentum sector scan."""
    intents, tickers, sectors = expand_query_intents("Which technology stocks are undervalued?")
    assert intents.get("valuation") is True
    assert "technology" in sectors


# ---------------------------------------------------------------------------
# v7: ambiguous-ticker security context, options/hedging routing, sector
# valuation screens, FX realism + dollar handling.
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("q,tok", [
    ("Hedge my LOW position against a market correction", "LOW"),
    ("Hedge my TGT position against a market correction", "TGT"),
    ("How likely is LOW to drop 25% in the next 3 months?", "LOW"),
    ("What options strategy makes sense for LOW before earnings?", "LOW"),
    ("What is the best put structure to protect TGT?", "TGT"),
    ("How does EWZ compare to TGT?", "TGT"),
    ("Should I own TGT or JPM?", "TGT"),
    ("Who are LOW's main competitors?", "LOW"),
    ("What does TGT do?", "TGT"),
])
def test_ambiguous_ticker_resolves_with_security_context(q, tok):
    """Ambiguous tokens (LOW/TGT are real tickers AND English words) must
    resolve as securities when the query uses stock/hedge/options/comparison
    framing."""
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


@pytest.mark.parametrize("q,tok", [
    ("AI infrastructure spending is booming", "AI"),
    ("low probability event", "LOW"),
    ("AI is a concept not a company", "AI"),
])
def test_ambiguous_ticker_stays_concept_without_security_context(q, tok):
    """The same ambiguous tokens must NOT resolve as securities in pure
    concept/prose contexts."""
    _i, tickers, _s = expand_query_intents(q)
    assert tok not in [str(t).upper() for t in tickers], (q, tickers)


def test_options_ask_about_named_stock_not_event_briefing():
    """'What options strategy makes sense for LOW before earnings?' must NOT
    be swallowed by the Current Events Briefing (earnings anchor) — the
    options/hedging path answers it with LOW in focus."""
    intents, tickers, _s = expand_query_intents(
        "What options strategy makes sense for LOW before earnings?")
    assert intents.get("options") is True
    assert "LOW" in [str(t).upper() for t in tickers]
    from financial_llm_engine import _is_focused_event_question
    assert _is_focused_event_question(
        "What options strategy makes sense for LOW before earnings?",
        intents, tickers) is False


def test_hedge_output_has_no_primary_literal():
    """The correlated-hedge line must interpolate the actual symbol, never
    print a literal '{primary}' placeholder."""
    from chatbot_eval.pipeline import mocked_pipeline
    from financial_llm_engine import generate_financial_analysis
    with mocked_pipeline():
        r = generate_financial_analysis(
            "Hedge my XOM position against a market correction") or ""
    assert "{primary}" not in r
    assert "XOM" in r


def test_sector_valuation_screen_engages_sector():
    """'Which technology stocks are undervalued?' produces a sector valuation
    screen that names the sector and multiple names."""
    from chatbot_eval.pipeline import mocked_pipeline
    from financial_llm_engine import generate_financial_analysis
    with mocked_pipeline():
        r = generate_financial_analysis("Which technology stocks are undervalued?") or ""
    assert "Valuation Screen" in r
    assert "technology" in r.lower()
    assert r.count("**") >= 6  # several ranked names


def test_fx_outlook_mentions_dollar_when_asked():
    """'Is the dollar strengthening?' must explicitly address the dollar
    (the eval rubric requires the word 'dollar' in the response)."""
    from chatbot_eval.pipeline import mocked_pipeline
    from financial_llm_engine import generate_financial_analysis
    with mocked_pipeline():
        r = generate_financial_analysis("Is the dollar strengthening?") or ""
    assert "dollar" in r.lower()
    assert "FX" in r


def test_fx_outlook_realistic_pair_levels():
    """The eval mock must produce realistic FX price levels (EURUSD ~1.0-1.3,
    never an arbitrary 200+ hash-derived price)."""
    from chatbot_eval.pipeline import _series_for
    import pandas as pd
    eur = _series_for("EURUSD=X", 40)["Close"]
    if isinstance(eur, pd.DataFrame):
        eur = eur.iloc[:, 0]
    px = float(eur.dropna().iloc[-1])
    assert 0.5 < px < 2.0, f"EURUSD mock price unrealistic: {px}"


# ---------------------------------------------------------------------------
# v8: sector display word-boundaries, macro/valuation security context,
# mega topic-pivot connectors ("In the same vein").
# ---------------------------------------------------------------------------


def test_sector_display_uses_canonical_name():
    """'Scan the technology sector' must label the scan 'Technology', not
    'Tech' (alias 'tech' is a substring of 'technology')."""
    from financial_llm_engine import _sector_display
    assert _sector_display("Scan the technology sector for opportunities",
                           "technology") == "Technology"
    assert _sector_display("Scan the healthcare sector", "healthcare") == "Healthcare"
    assert _sector_display("Give me tech stocks", "technology") == "Tech"


@pytest.mark.parametrize("q,tok", [
    ("What happens to TGT and gold if the Fed cuts rates", "TGT"),
    ("Is TGT overvalued or undervalued at current levels", "TGT"),
    ("How likely is COST to fall to $837 in 18 months?", "COST"),
])
def test_macro_and_valuation_ambiguous_ticker_context(q, tok):
    """Ambiguous tickers in macro/transmission and valuation framing must
    resolve as securities."""
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


def test_mega_split_on_in_the_same_vein():
    """'... ? In the same vein, should I own bonds?' is a second distinct
    task and must be split, not merged into the first part."""
    from financial_llm_engine import _decompose_mega_query
    parts = _decompose_mega_query(
        "is RIVN overvalued or undervalued?. In the same vein, should i own "
        "long-duration bonds right now?") or []
    assert len(parts) >= 2, parts
    assert any("bond" in p for p in parts), parts


def test_mega_split_on_additionally():
    """'... . Additionally, should I own bonds?' splits into two tasks."""
    from financial_llm_engine import _decompose_mega_query
    parts = _decompose_mega_query(
        "what are the best semiconductors stocks to buy now?. Additionally, "
        "should i own long-duration bonds right now?") or []
    assert len(parts) >= 2, parts
    assert any("bond" in p for p in parts), parts


# ── v9: news/moat/dividend/options/positioning context + single-letter and ──
# ── pair-coordination resolution ─────────────────────────────────────────────
@pytest.mark.parametrize("q, tok", [
    ("What happened to COST this week?", "COST"),
    ("Is the news flow bullish or bearish for COST?", "COST"),
    ("What is the latest news on COST?", "COST"),
    ("Does LOW pay a dividend?", "LOW"),
    ("Has LOW been growing its dividend?", "LOW"),
    ("What is the dividend yield of LOW?", "LOW"),
    ("How should I size positions in LOW?", "LOW"),
    ("Should I sell covered calls on TGT?", "TGT"),
    ("Is a bull call spread on TGT a good idea?", "TGT"),
    ("How does TGT compare to AAPL?", "TGT"),
    ("Should I own BA or TGT?", "TGT"),
    ("Analyze TGT's competitive moat versus AAPL: pricing power", "TGT"),
    ("how likely is F to fall to $33 in 6 months", "F"),
    ("how likely is C to fall to $110 in 12 months", "C"),
    ("compare XLF and TGT what is the outlook for MSFT?", "TGT"),
])
def test_v9_ambiguous_ticker_context(q, tok):
    """News-recap, moat, dividend, options, positioning and single-letter /
    pair-coordination phrasings must all resolve the named security."""
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


@pytest.mark.parametrize("q, toks", [
    # Concepts must NEVER resolve as tickers
    ("what is the net loss for the quarter", []),
    ("probabilities that sum to 100%", []),
    ("key drivers of growth", []),
    ("base case scenario", []),
    ("AI infrastructure spending", []),
    ("competitive moat and switching costs", []),
    ("target price of 50", []),
    ("what do they cost", []),
    ("low volatility stocks", []),
    ("grade was a C", []),
    ("plan B is better", []),
    ("the F word", []),
])
def test_v9_no_concept_leak(q, toks):
    """The new single-letter and coordination rules must not resolve concepts
    (MOAT, TGT-as-target, COST-as-cost, C-as-grade, ...) as securities."""
    _i, tickers, _s = expand_query_intents(q)
    assert set(tickers) == set(toks), (q, tickers)


def test_v9_single_letter_no_false_positive_in_mega():
    """The NVDA deep-dive prompt (which contains 'P/E and FCF yield' and
    'durable moat') must not resolve a phantom single-letter 'E' ticker."""
    q = ("Conduct a full fundamental, quantitative, macroeconomic and "
         "sentiment-driven investment assessment of NVIDIA (NVDA). Reconstruct "
         "revenue and earnings drivers by segment. Evaluate competitive "
         "position against AMD, custom ASICs and hyperscaler silicon. "
         "Determine whether CUDA represents a durable moat. Build a DCF using "
         "at least three explicit scenarios. Perform a reverse DCF. Compare "
         "using EV/Revenue, EV/EBITDA, P/E and FCF yield. Analyze sensitivity "
         "to Fed policy, real interest rates, USD strength and the "
         "semiconductor cycle. Construct at least five scenarios and assign "
         "probabilities. Identify the 10 strongest arguments. Give me "
         "probability-weighted fair value and the probability of a >30% "
         "drawdown.")
    _i, tickers, _s = expand_query_intents(q)
    assert set(tickers) == {"NVDA", "AMD"}, tickers


def test_v9_mega_split_unpunctuated_question_boundary():
    """'compare XLF and TGT what is the outlook for MSFT?' has NO punctuation
    between clauses — the unpunctuated-boundary pass must still split it."""
    from financial_llm_engine import _decompose_mega_query
    parts = _decompose_mega_query(
        "compare XLF and TGT what is the outlook for MSFT?") or []
    assert len(parts) >= 2, parts
    joined = " ".join(parts).lower()
    assert "compare" in joined and "outlook" in joined, parts


def test_v9_statement_not_split_by_copula():
    """Copula statements ('NVDA is a good buy') must never be split into two
    fake mega-parts."""
    from financial_llm_engine import _decompose_mega_query
    for q in ("NVDA is a good buy", "AMD can rally 20%", "COST was up this week"):
        assert _decompose_mega_query(q) is None, q


# ── v10: drawdown/straddle/sector/portfolio context + task-word mega ────────
# ── boundaries + identical-profile mega split + None-price crash guard ───────
@pytest.mark.parametrize("q, tok", [
    ("What are the chances of a drawdown in LOW?", "LOW"),
    ("What are the chances of a drawdown in TGT?", "TGT"),
    ("Would a straddle work on TGT?", "TGT"),
    ("What sector is TGT in?", "TGT"),
    ("How should I think about LOW in my portfolio?", "LOW"),
    ("What is the probability of a recession in the next 12 months, and how "
     "should I position TGT defensively?", "TGT"),
])
def test_v10_ambiguous_ticker_context(q, tok):
    """Drawdown, straddle, sector and portfolio phrasings resolve the ticker."""
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


@pytest.mark.parametrize("q, checks", [
    ("give me a trade setup for XLI with entry, stop and target what is the "
     "outlook for the 10-year treasury yield?",
     [lambda p: "trade setup" in p.lower(), lambda p: "10-year" in p.lower()]),
    ("give me a trade setup for ADBE with entry, stop and target which is the "
     "better investment, C or WFC?",
     [lambda p: "trade setup" in p.lower(), lambda p: "C or WFC" in p]),
    ("what is the outlook for the yield curve? Now: what is the outlook for "
     "the 10-year treasury yield?",
     [lambda p: "yield curve" in p.lower(), lambda p: "10-year" in p.lower()]),
    ("should i own long-duration bonds right now?. On a different note, what "
     "is the outlook for quantitative tightening?",
     [lambda p: "bonds" in p.lower(), lambda p: "quantitative tightening" in p.lower()]),
    ("what is the correlation between platinum and the S&P 500?. Separately, "
     "which is the better investment, SPY or TGT?",
     [lambda p: "correlation" in p.lower(), lambda p: "SPY or TGT" in p]),
])
def test_v10_mega_task_word_and_connector_splits(q, checks):
    """Setup keywords before a question word ('target what'), 'On a different
    note', 'Separately' and identical-profile-but-distinct asks must all split
    into their own mega-parts."""
    from financial_llm_engine import _decompose_mega_query
    parts = _decompose_mega_query(q) or []
    assert len(parts) >= 2, (q, parts)
    joined = " ".join(parts)
    for chk in checks:
        assert any(chk(p) for p in parts), (q, joined)


def test_v10_identical_profile_distinct_asks_split():
    """Two same-intent asks about different topics stay separate parts."""
    from financial_llm_engine import _decompose_mega_query
    parts = _decompose_mega_query(
        "what is the outlook for the yield curve? Now: what is the outlook for "
        "the 10-year treasury yield?") or []
    assert len(parts) >= 2, parts


def test_v10_mega_part_crash_is_contained():
    """A crashing per-part builder (None price) must not kill the whole
    multi-part answer — the fallback section still appears."""
    from financial_llm_engine import _build_mega_response, _decompose_mega_query
    parts = _decompose_mega_query(
        "Analyze the technicals for UAL, then if the trend is constructive give "
        "me a trade setup with entry, stop and target, and finally recommend "
        "how to hedge the position.") or []
    assert len(parts) >= 3, parts
    # The hedge part is last; the builder must render all parts or degrade.
    out = _build_mega_response(" ".join(parts), parts, None)
    assert out is None or "Multi-Part Analysis" in out, out


def test_v11_play_and_protect_gains_extract_tickers():
    """'best way to play TGT' / 'protect gains in TGT' must resolve TGT/LOW
    as securities (verb+preposition framing), not drop them."""
    from financial_llm_engine import expand_query_intents
    for q, want in [
        ("What is the best way to play TGT this week?", ["TGT"]),
        ("What is the best way to protect gains in TGT?", ["TGT"]),
        ("What is the best way to protect gains in LOW?", ["LOW"]),
        ("How to play LOW into earnings?", ["LOW"]),
    ]:
        _, t, _ = expand_query_intents(q)
        assert sorted(t) == sorted(want), (q, t)


def test_v11_play_phrase_no_false_positive():
    """'play the market' / plain prose never resolves an ambiguous ticker."""
    from financial_llm_engine import _is_equity_reference
    assert not _is_equity_reference("What is the best way to play the market this week?", "TGT")
    assert not _is_equity_reference("gains in the sector are broad", "TGT")


def test_v11_mega_imperative_boundary_split():
    """'...target compare JNJ and GOOG give me the key support' must split into
    setup / comparison / technical parts without punctuation."""
    from financial_llm_engine import _decompose_mega_query, expand_query_intents
    parts = _decompose_mega_query(
        "what is the probability UBER reaches $41 in 24 months?. Meanwhile, give "
        "me a trade setup for BIIB with entry, stop and target compare JNJ and "
        "GOOG give me the key support and resistance") or []
    assert len(parts) >= 4, parts
    for p in parts:
        _, t, i = expand_query_intents(p)
        pl = p.lower()
        if "probability" in pl:
            assert "UBER" in t, (p, t)
        if "setup" in pl:
            assert "BIIB" in t, (p, t)
        if "compare" in pl:
            assert {"JNJ", "GOOG"} <= set(t), (p, t)
        if "key support" in pl:
            assert any(w in pl for w in ("support", "resistance")), p


def test_v11_mega_imperative_boundary_no_false_split():
    """Task words followed by plain prose never create bogus mega-parts."""
    from financial_llm_engine import _decompose_mega_query
    for q in [
        "entry, stop and target are the key levels to watch",
        "what is the best way to play the market this week?",
    ]:
        parts = _decompose_mega_query(q)
        assert parts is None, (q, parts)


def test_v11_mega_two_macro_parts_both_kept():
    """Two genuinely different macro asks with the same label must both be
    answered — the label de-dup only drops near-duplicate prose."""
    from financial_llm_engine import _decompose_mega_query, _build_mega_response
    q = "should i own long-duration bonds right now?. On a different note, what is the outlook for quantitative tightening?"
    parts = _decompose_mega_query(q) or []
    assert len(parts) >= 2, parts
    out = _build_mega_response(q, parts, None) or ""
    assert out.count("### Part ") >= 2, out
    assert "quantitative tightening" in out.lower(), out
    # The bonds part renders its macro grounding under the "Duration" topic
    # label (2s10s / long-duration hedges) — either phrasing proves it landed.
    assert "bonds" in out.lower() or "duration" in out.lower(), out


def test_v11_mega_near_duplicate_parts_still_deduped():
    """A re-asked/near-duplicate part under the same label IS still dropped."""
    from financial_llm_engine import _decompose_mega_query, _build_mega_response
    parts = [
        "what is the outlook for the economy right now",
        "what is the outlook for the economy",
    ]
    out = _build_mega_response(" ".join(parts), parts, None) or ""
    # The second part is a strict sub-question of the first (containment), so
    # the de-dup drops it as a re-ask while keeping the fuller version.
    assert out.count("### Part ") == 1, out


def test_v11_mega_distinct_instruments_same_label_both_kept():
    """Two 'outlook for X' asks under the same 'Analysis' label but with
    DIFFERENT instruments (VIX vs PYPL, GILD vs VIX, LI vs Nasdaq) must both
    be answered — template text overlap must never drop a distinct ask."""
    from financial_llm_engine import _decompose_mega_query, _build_mega_response
    for q in [
        "what is the outlook for the VIX?. Next, what is the outlook for PYPL?",
        "what is the outlook for LI?. And since we are on the topic, what is the outlook for the Nasdaq?",
        "what is the outlook for GILD?. Meanwhile, what is the outlook for the VIX?",
    ]:
        parts = _decompose_mega_query(q) or []
        assert len(parts) >= 2, (q, parts)
        out = _build_mega_response(q, parts, None) or ""
        assert out.count("### Part ") >= 2, (q, out)


def test_v11_mega_same_instrument_deduped():
    """Two asks about the SAME instrument under one label are a re-ask and
    must collapse to a single part."""
    from financial_llm_engine import _build_mega_response
    parts = ["what is the outlook for AAPL", "what is the outlook for AAPL next quarter"]
    out = _build_mega_response(" ".join(parts), parts, None) or ""
    assert out.count("### Part ") == 1, out


# ---------------------------------------------------------------------------
# v12: stopword tickers (ARE/ALL/AM/CAN/...) resolve only when UPPERCASE +
# explicit security context; prose plurals and verb forms never do.
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("q,tok", [
    ("buy ARE", "ARE"),
    ("ARE stock analysis", "ARE"),
    ("outlook for ARE", "ARE"),
    ("Is ARE undervalued at current levels", "ARE"),
    ("Allstate ALL stock", "ALL"),
    ("buy ALL shares", "ALL"),
    ("ALL stock vs TGT", "ALL"),
    ("play RUN this week", "RUN"),
    ("ARE or NVDA which is better", "ARE"),
])
def test_v12_stopword_ticker_resolves_with_uppercase_security_context(q, tok):
    """Genuine universe tickers that are also core English words (ARE, ALL,
    RUN, ...) must resolve when written in UPPERCASE with explicit security
    context — previously they were unconditionally dropped by _STOPWORDS."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


@pytest.mark.parametrize("q", [
    "the are of investing",
    "all stocks are down",
    "can you analyze the market",
    "how are stocks doing today",
    "how much are prices going up",
    "Which technology stocks are undervalued?",
    "real estate is expensive",
    "buy are",          # lowercase ticker = prose, per case-sensitivity rule
    "there are no good buys today",
])
def test_v12_stopword_ticker_never_leaks_from_prose(q):
    """Prose uses of stopword tickers ('are', 'all', 'can') must never resolve
    — even in equity-adjacent phrasing like 'stocks are undervalued' or
    'all stocks' (plural). Lowercase tickers are prose by convention."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents(q)
    leaked = [str(t) for t in tickers if str(t).upper() in ("ARE", "ALL", "AM", "CAN")]
    assert not leaked, (q, tickers)


def test_v12_gold_remains_commodity_not_equity():
    """'gold' (lowercase) must stay a commodity (GC=F) and never resolve as
    the GOLD equity ticker via the stopword escape."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents("What happens to TGT and gold if the Fed cuts rates")
    uppers = [str(t).upper() for t in tickers]
    assert "TGT" in uppers, tickers
    assert "GOLD" not in uppers, tickers
    assert any("GC=F" == str(t) for t in tickers), tickers


@pytest.mark.parametrize("q,tok", [
    ("Give me a trade setup for ARE with entry, stop and target.", "ARE"),
    ("Give me a trade setup for ALL with entry, stop and target.", "ALL"),
    ("trade setup for CAN with entry and stop", "CAN"),
])
def test_v12_trade_setup_phrase_resolves_stopword_ticker(q, tok):
    """'trade setup for X' is unambiguous security context — stopword tickers
    (ARE/ALL/CAN) must resolve in it (this closed the last 13/100k stress
    failures)."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)


def test_v12_trade_setup_prose_never_resolves():
    """'the setup for all users is ready' is prose — lowercase 'all' must not
    resolve even with the trade-setup phrase present."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents("the setup for all users is ready")
    assert not [str(t) for t in tickers if str(t).upper() == "ALL"], tickers


# ---------------------------------------------------------------------------
# v12c: FX named-pair asks must lead with the specific pair (not just a
# generic G10 momentum ranking), with an evidence-based Fed-transmission note.
# ---------------------------------------------------------------------------

def test_fx_named_pair_leads_with_pair_and_fed_transmission():
    """'How will USD/BRL react to the Fed?' must lead with the USDBRL quote and
    a Fed-transmission note, not only the generic momentum ranking."""
    from financial_llm_engine import generate_financial_analysis
    from chatbot_eval.pipeline import mocked_pipeline
    with mocked_pipeline():
        r = generate_financial_analysis("How will USD/BRL react to the Fed?") or ""
        assert "US Dollar / Brazilian Real" in r, r[:300]
        assert "Fed transmission" in r, r[:300]


def test_fx_generic_ask_still_ranks_pairs():
    """A generic currency ask keeps the momentum ranking and does not
    fabricate a named pair."""
    from financial_llm_engine import generate_financial_analysis
    from chatbot_eval.pipeline import mocked_pipeline
    with mocked_pipeline():
        r = generate_financial_analysis("Which currency is the best buy right now?") or ""
        assert "Euro / US Dollar" in r or "US Dollar / Yen" in r, r[:300]


# ---------------------------------------------------------------------------
# v12d: 'setup for X' phrases must not leak into the ambiguous-token path —
# prose like 'the setup for key support levels' must never resolve KEY/AI/SUM.
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("q", [
    "what is the setup for key support levels",
    "the setup for AI infrastructure spend",
    "setup for SUM of the parts",
    "the setup for all users is ready",
])
def test_v12_setup_prose_never_resolves_ambiguous_tokens(q):
    """'setup for X' is security context ONLY on the strict stopword path;
    ambiguous concept tokens (KEY/AI/SUM/ALL) must never resolve from prose."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents(q)
    leaked = [str(t) for t in tickers if str(t).upper() in ("KEY", "AI", "SUM", "ALL")]
    assert not leaked, (q, tickers)


@pytest.mark.parametrize("q,tok", [
    ("Give me a trade setup for KEY stock", "KEY"),
    ("Give me a trade setup for AI stock", "AI"),
])
def test_v12_ambiguous_ticker_still_resolves_with_own_phrase(q, tok):
    """Ambiguous tokens keep resolving via their own security phrases
    ('KEY stock', 'AI stock') — the setup gate must not break those."""
    from financial_llm_engine import expand_query_intents
    _i, tickers, _s = expand_query_intents(q)
    assert tok in [str(t).upper() for t in tickers], (q, tickers)
